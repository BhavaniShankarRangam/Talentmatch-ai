"""Resume text extraction into referenceable segments (page + section + line)."""
import io
import re
from dataclasses import asdict, dataclass, field

HEADINGS: dict[str, list[str]] = {
    "summary": ["summary", "profile", "professional summary", "about me", "objective"],
    "experience": [
        "experience", "professional experience", "work experience", "employment history",
        "employment", "work history",
    ],
    "skills": ["skills", "technical skills", "core skills", "key skills", "skills & tools"],
    "projects": ["projects", "selected projects", "project experience", "key projects"],
    "education": ["education", "education & training"],
    "certifications": [
        "certifications", "certificates", "certification", "licenses & certifications",
        "certifications & licenses",
    ],
    "personal": ["personal details", "personal information", "personal data"],
}
_HEADING_LOOKUP = {alias: key for key, aliases in HEADINGS.items() for alias in aliases}

MIN_USABLE_CHARS = 80


class ParseError(Exception):
    pass


@dataclass
class Segment:
    index: int
    page: int | None
    section: str
    text: str

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict) -> "Segment":
        return cls(index=d["index"], page=d.get("page"), section=d["section"], text=d["text"])


@dataclass
class ParseResult:
    status: str  # parsed | needs_review | failed
    segments: list[Segment] = field(default_factory=list)
    page_count: int | None = None
    char_count: int = 0
    confidence: str = "none"  # high | medium | low | none
    notes: list[str] = field(default_factory=list)
    error: str | None = None


def _normalize_heading(line: str) -> str:
    return re.sub(r"[^a-z& ]", "", line.lower()).strip()


def _to_segments(lines: list[tuple[int | None, str]]) -> list[Segment]:
    section = "header"
    segments: list[Segment] = []
    for page, raw in lines:
        line = " ".join(raw.split())
        if not line:
            continue
        key = _HEADING_LOOKUP.get(_normalize_heading(line)) if len(line) <= 40 else None
        if key:
            section = key
            continue
        segments.append(Segment(index=len(segments), page=page, section=section, text=line[:1000]))
    return segments


def _pdf_lines(data: bytes, max_pages: int) -> tuple[list[tuple[int | None, str]], int]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ParseError("PDF is password-protected and cannot be read.")
    n = len(reader.pages)
    if n > max_pages:
        raise ParseError(f"PDF has {n} pages; the limit is {max_pages}.")
    lines: list[tuple[int | None, str]] = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        lines.extend((i, ln) for ln in text.splitlines())
    return lines, n


def _docx_lines(data: bytes) -> list[tuple[int | None, str]]:
    import docx

    d = docx.Document(io.BytesIO(data))
    lines: list[tuple[int | None, str]] = [(None, p.text) for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                lines.extend((None, p.text) for p in cell.paragraphs)
    return lines


def parse_document(data: bytes, ext: str, max_pages: int = 40) -> ParseResult:
    page_count: int | None = None
    try:
        if ext == ".pdf":
            lines, page_count = _pdf_lines(data, max_pages)
        elif ext == ".docx":
            lines = _docx_lines(data)
        else:
            return ParseResult(status="failed", error=f"Unsupported document type {ext}.")
    except ParseError as e:
        return ParseResult(status="failed", error=str(e))
    except Exception as e:  # malformed/corrupt documents raise many different parser errors
        return ParseResult(
            status="failed",
            error=f"Document could not be read ({type(e).__name__}); it may be corrupt or malformed.",
        )

    segments = _to_segments(lines)
    chars = sum(len(s.text) for s in segments)
    sections = {s.section for s in segments} - {"header"}
    notes: list[str] = []

    if chars < MIN_USABLE_CHARS:
        notes.append(
            "Little or no extractable text was found. The file may be scanned or image-only; "
            "OCR is not enabled in this milestone."
        )
        return ParseResult(
            status="needs_review", segments=segments, page_count=page_count, char_count=chars,
            confidence="none", notes=notes,
        )

    if chars >= 600 and len(sections) >= 3:
        confidence = "high"
    elif chars >= 250 and len(sections) >= 1:
        confidence = "medium"
    else:
        confidence = "low"
    if not sections:
        notes.append("No standard resume section headings detected; section references may be imprecise.")
    if confidence == "low":
        notes.append("Limited text extracted; some evidence may have been missed.")
    return ParseResult(
        status="parsed", segments=segments, page_count=page_count, char_count=chars,
        confidence=confidence, notes=notes,
    )


def extract_raw_text(data: bytes, ext: str, max_pages: int = 40) -> str:
    """Plain text with original line structure (used for job description uploads)."""
    try:
        lines = _pdf_lines(data, max_pages)[0] if ext == ".pdf" else _docx_lines(data)
    except ParseError:
        raise
    except Exception as e:
        raise ParseError(f"Document could not be read ({type(e).__name__}).")
    return "\n".join(t for _, t in lines if t.strip())
